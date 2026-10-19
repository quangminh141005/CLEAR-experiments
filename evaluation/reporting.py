"""Vietnamese research report from recorded runs only; no inference or network."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd

from .comparison import PAPER_URL, SUPPLEMENT_URL, compare, references
from .storage import file_hash


SECTIONS = ["Tóm tắt kết quả", "Các công việc đã thực hiện", "Môi trường thực nghiệm",
    "Mô tả dữ liệu Indiana, PadChest và CheXpert", "Quy trình inference của CLEAR", "Phương pháp đánh giá",
    "Kết quả trên Indiana", "Kết quả trên PadChest", "Kết quả trên CheXpert", "So sánh với CLEAR paper",
    "Phân tích False Positive và False Negative", "Những hạn chế của quá trình reproduction",
    "Kết luận và hướng nghiên cứu tiếp theo", "Phụ lục và tài liệu tham khảo"]
DATASETS = ("indiana", "padchest", "chexpert")


def fmt(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "Không xác định"
    return f"{value:.3f}" if isinstance(value, float) else str(value)


def ci_text(bounds):
    if not bounds or bounds.get("low") is None:
        return "Không xác định"
    return f"[{fmt(bounds['low'])}; {fmt(bounds['high'])}]"


class Report:
    def __init__(self):
        self.blocks = []

    def heading(self, text, level=1):
        self.blocks.append(("heading", text, level))

    def paragraph(self, text):
        self.blocks.append(("paragraph", text))

    def table(self, headers, rows):
        self.blocks.append(("table", headers, [[fmt(v) for v in row] for row in rows]))

    def image(self, path, caption):
        self.blocks.append(("image", Path(path).resolve(), caption))

    def markdown(self, output):
        import os
        parts = []
        for kind, *data in self.blocks:
            if kind == "heading":
                parts.append("#" * data[1] + " " + data[0])
            elif kind == "paragraph":
                parts.append(data[0])
            elif kind == "table":
                def escape(value):
                    return str(value).replace("|", "\\|").replace("\n", " ")
                parts.append("| " + " | ".join(map(escape, data[0])) + " |\n| " + " | ".join("---" for _ in data[0]) + " |\n" +
                             "\n".join("| " + " | ".join(map(escape, row)) + " |" for row in data[1]))
            elif kind == "image":
                relative = os.path.relpath(data[0], Path(output).parent)
                parts.append(f"![{data[1]}](<{relative}>)\n\n{data[1]}")
        return "\n\n".join(parts) + "\n"

    def docx(self, output, template=None):
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor

        document = Document(template) if template else Document()
        if template:
            # A template supplies styles and page settings; remove placeholder body content.
            body = document._element.body
            for element in list(body):
                if element.tag != qn("w:sectPr"):
                    body.remove(element)
        for section in document.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.left_margin = section.right_margin = Cm(2)
            section.top_margin = section.bottom_margin = Cm(2)
        normal = document.styles["Normal"]
        normal.font.name, normal.font.size = "Arial", Pt(10)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for kind, *data in self.blocks:
            if kind == "heading":
                document.add_paragraph(data[0], style="Title" if data[1] == 1 else f"Heading {min(data[1] - 1, 3)}")
            elif kind == "paragraph":
                document.add_paragraph(data[0])
            elif kind == "table":
                headers, rows = data
                table = document.add_table(rows=1, cols=len(headers))
                table.autofit = False
                # Long finding names receive more space than numeric columns.
                widths = [5.0] + [12.0 / (len(headers) - 1)] * (len(headers) - 1) if len(headers) > 1 else [17.0]
                for column, width in zip(table.columns, widths):
                    column.width = Cm(width)
                for cell, width, text in zip(table.rows[0].cells, widths, headers):
                    cell.width, cell.text = Cm(width), text
                for row in rows:
                    for cell, width, text in zip(table.add_row().cells, widths, row):
                        cell.width, cell.text = Cm(width), text
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for index, row in enumerate(table.rows):
                    for column, cell in enumerate(row.cells):
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        properties = cell._tc.get_or_add_tcPr()
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E3EDF4" if index == 0 else ("F5F7F9" if index % 2 == 0 else "FFFFFF"))
                        properties.append(shade)
                        borders = OxmlElement("w:tcBorders")
                        for edge in ("top", "left", "bottom", "right"):
                            border = OxmlElement(f"w:{edge}")
                            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                                border.set(qn(f"w:{key}"), value)
                            borders.append(border)
                        properties.append(borders)
                        margins = OxmlElement("w:tcMar")
                        for edge in ("top", "left", "bottom", "right"):
                            item = OxmlElement(f"w:{edge}")
                            item.set(qn("w:w"), "80")
                            item.set(qn("w:type"), "dxa")
                            margins.append(item)
                        properties.append(margins)
                        for paragraph in cell.paragraphs:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
                            paragraph.paragraph_format.space_after = Pt(3)
                            for run in paragraph.runs:
                                run.font.size = Pt(8)
                                run.bold = index == 0
                document.add_paragraph()
            elif kind == "image":
                # Scale height as well as width so long per-class charts fit A4.
                from PIL import Image
                with Image.open(data[0]) as source:
                    width, height = source.size
                inches_w = min(6.6, 8.0 * width / height)
                document.add_picture(str(data[0]), width=Cm(inches_w * 2.54))
                document.add_paragraph(data[1], style="Caption")
        footer = document.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("CLEAR · Báo cáo nghiên cứu · ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        document.core_properties.title = "Đánh giá khả năng tái lập mô hình CLEAR"
        document.core_properties.subject = "Đánh giá zero shot trên Indiana PadChest và CheXpert"
        document.save(output)


def read_runs(paths):
    runs = {}
    for name, location in paths.items():
        root = Path(location).resolve()
        run_path = root / "run.json"
        if not run_path.is_file():
            raise FileNotFoundError(f"No pipeline run.json at {root}; legacy experiments are not silently imported")
        run = json.loads(run_path.read_text())
        if run.get("schema_version") != 1 or run["dataset"] != name:
            raise ValueError(f"Run schema/dataset mismatch: {root}")
        entry = {"root": root, "run": run, "metrics": None}
        if run["status"] == "completed":
            entry["metrics"] = json.loads((root / "metrics.json").read_text())
            if entry["metrics"]["images"] != len(run["cohort"]["records"]):
                raise ValueError(f"Metric/cohort size mismatch in {root}")
            digest = file_hash(root / "predictions.json")
            if digest != run.get("predictions_sha256") or digest != entry["metrics"].get("predictions_sha256"):
                raise ValueError(f"Predictions/metrics provenance mismatch in {root}")
        runs[name] = entry
    return runs


def build_report(runs):
    report = Report()
    report.heading("Đánh giá khả năng tái lập mô hình CLEAR")
    report.paragraph("Báo cáo đánh giá zero shot trên Indiana, PadChest sample và CheXpert. "
        "Kết quả thực nghiệm chỉ được lấy từ các lần chạy đã hoàn tất do người dùng cung cấp. "
        "Số liệu trích dẫn từ bài báo được ghi riêng; dữ liệu thiếu được ghi là chưa thực hiện hoặc không xác định.")
    for number, title in enumerate(SECTIONS, 1):
        report.heading(f"{number} {title}", 2)
        if number == 1:
            rows = []
            for dataset in DATASETS:
                entry = runs.get(dataset)
                metric = entry["metrics"] if entry else None
                rows.append([dataset, "Đã thực hiện" if metric else "Chưa thực hiện" if not entry else f"Chưa hoàn tất ({entry['run']['status']})",
                    metric["images"] if metric else None, metric["macro_auroc"] if metric else None,
                    ci_text(metric["macro_auroc_ci95"]) if metric else "Chưa thực hiện"])
            report.table(["Dữ liệu", "Trạng thái", "Số ảnh", "Macro AUROC", "KTC 95%"], rows)
            report.paragraph("Macro AUROC chỉ tính trên các lớp có cả nhãn dương và âm. Các bộ lớp hợp lệ có thể khác nhau; "
                "không suy ra thứ hạng giữa các tập dữ liệu từ macro đơn lẻ.")
        elif number == 2:
            report.paragraph("Mã nguồn đã tích hợp bộ nạp dữ liệu, chính sách nhãn, tiền xử lý CLEAR, chiếu concept "
                "vào không gian SFR, baseline trực tiếp, lưu và tiếp tục dự đoán, thống kê bootstrap, "
                "đối chiếu số liệu công bố và xuất báo cáo. Việc có mã nguồn không đồng nghĩa các thí nghiệm đã được chạy.")
        elif number == 3:
            if not runs:
                report.paragraph("Chưa thực hiện. Chưa có thông tin môi trường thực nghiệm từ máy chủ.")
            for dataset, entry in runs.items():
                env = entry["run"].get("environment", {})
                report.heading(dataset, 3)
                report.table(["Thuộc tính", "Giá trị"], [["Python", env.get("python", "Chưa ghi nhận")],
                    ["Nền tảng", env.get("platform", "Chưa ghi nhận")], ["GPU", ", ".join(env.get("gpu_names", [])) or "Không ghi nhận"],
                    ["CUDA trong wheel", env.get("cuda_runtime")],
                    *[[name, version] for name, version in env.get("packages", {}).items()]])
            report.paragraph("Tesla K80 có compute capability 3.7. Profile legacy dùng Python 3.10, PyTorch 1.12.1/cu102, "
                "FP32 và attention thông thường; khả năng hoạt động thực tế phụ thuộc driver và wheel của máy chủ. "
                "Prompt SFR 7B được chuẩn bị riêng trên CPU đủ RAM hoặc GPU tương thích.")
        elif number == 4:
            report.paragraph("Indiana sử dụng Problems ở cấp báo cáo và ghép với ảnh qua uid; Pulmonary Atelectasis được ánh xạ "
                "thành atelectasis. uid chỉ là nhóm nghiên cứu thay thế nếu không có bản đồ bệnh nhân xác minh. "
                "PadChest sample ghép ImageID với Labels và PatientID; không đại diện toàn bộ benchmark bác sĩ gán nhãn. "
                "CheXpert lấy một tệp annotation chỉ định, giữ nhãn bất định -1 và nhãn trống theo chính sách đã ghi nhận.")
            for dataset, entry in runs.items():
                run = entry["run"]
                metadata = run["cohort"]["metadata"]
                report.heading(dataset, 3)
                report.table(["Thuộc tính", "Giá trị"], [[key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value]
                    for key, value in metadata.items()])
                report.paragraph("Chính sách nhãn: " + "; ".join(f"{key}={run['configuration'].get(key)}" for key in
                    ("unmentioned_policy", "normal_policy", "uncertain_policy", "missing_policy", "image_policy", "view")))
        elif number == 5:
            report.paragraph("Theo bài báo, ngân hàng concept được tạo bằng cách trích xuất observations từ 227.835 báo cáo, "
                "khử trùng lặp và chuẩn hóa chữ thường thành 368.294 observations. Pipeline sử dụng nguyên assets phát hành, "
                "không trích concept mới từ báo cáo đánh giá.")
            report.paragraph("Ảnh được resize bicubic về 448 × 448, chuyển RGB và chuẩn hóa bằng mean 101.48761, "
                "std 83.43944 trên thang 0–255. Encoder DINOv2 ViT B/14 với registers của checkpoint CLEAR tạo vector 768 chiều. "
                "Cosine với toàn bộ 368.294 observations tạo vector s; nhân s với ma trận SFR 368.294 × 4.096 rồi chuẩn hóa L2. "
                "Hai prompt dương và âm được mã hóa bằng SFR trong cùng không gian; softmax trên hai cosine tạo điểm bệnh. "
                "Thứ tự concept được giữ nguyên. Xử lý theo chunk chỉ giới hạn bộ nhớ, không cắt top k hoặc loại concept.")
            report.paragraph("Baseline direct dùng text encoder CLEAR trong không gian 768 chiều và không thực hiện phép chiếu SFR. "
                "Điểm softmax không phải xác suất lâm sàng đã hiệu chỉnh. Mã get_embed.py phát hành dùng văn bản raw cho nhánh SFR, "
                "trong khi bài báo mô tả instruction; recipe và revision prompt cache được lưu để kiểm tra sai khác này.")
            for dataset, entry in runs.items():
                run = entry["run"]
                report.paragraph(f"{dataset}: inference_mode={run['inference_mode']}; assets revision={run.get('assets', {}).get('revision')}; "
                    f"prompt recipe={json.dumps((run.get('prompts') or {}).get('metadata'), ensure_ascii=False)}.")
        elif number == 6:
            report.paragraph("AUROC là chỉ số chính, tính riêng từng lớp trên nhãn xác định; lớp thiếu dương hoặc âm nhận null. "
                "AP không xác định khi không có mẫu dương. Precision, recall, specificity và F1 dùng ngưỡng ghi trong metrics; "
                "mẫu thiếu nhãn không được đưa vào confusion matrix. Bootstrap percentile 95% lấy lại mẫu theo patient_id, "
                "giữ toàn bộ ảnh cùng bệnh nhân; nếu chỉ có study proxy thì độ bất định chỉ phản ánh nhóm proxy đó. "
                "Số bootstrap hợp lệ được lưu, lớp của macro được giữ cố định giữa các lần lấy mẫu.")
            report.paragraph("Ngưỡng cố định mặc định là 0.5. Ngưỡng chọn theo maximum F1 chỉ hợp lệ khi được ước lượng trên "
                "validation độc lập, có patient IDs không giao với evaluation. Không chọn ngưỡng trên tập test. "
                "Tập validation/test mới được chia theo hash patient ID; subset giới hạn số nhóm và giữ toàn bộ ảnh của nhóm.")
            for dataset, entry in runs.items():
                if entry["metrics"]:
                    metric = entry["metrics"]
                    policy = metric["threshold_policy"]
                    description = (f"cố định {policy.get('value')}" if policy["type"] == "fixed" else
                        f"chọn trên validation độc lập theo {policy['provenance']['selection']}; "
                        f"prediction SHA256={policy['provenance']['validation_predictions_sha256']}")
                    report.paragraph(f"{dataset}: số bootstrap={metric['bootstrap_replicates']}; "
                        f"chính sách ngưỡng={description}; "
                        f"bootstrap macro hợp lệ={metric['macro_auroc_ci95']['valid_resamples']}.")
        elif number in (7, 8, 9):
            dataset = DATASETS[number - 7]
            entry = runs.get(dataset)
            if not entry or not entry["metrics"]:
                report.paragraph("Chưa thực hiện hoặc chưa hoàn tất. Không có số liệu thực nghiệm hợp lệ để báo cáo.")
                continue
            metric = entry["metrics"]
            report.paragraph(f"Nguồn kết quả thực nghiệm: {entry['root']}. Mode {entry['run']['inference_mode']}. "
                f"Macro AUROC={fmt(metric['macro_auroc'])}, KTC 95%={ci_text(metric['macro_auroc_ci95'])}; "
                f"{len(metric['valid_auroc_classes'])}/{metric['number_of_classes']} lớp có AUROC xác định.")
            report.table(["Phát hiện", "AUROC", "KTC 95%", "AP", "F1"], [[label, values["auroc"],
                ci_text(values["ci95"]["auroc"]), values["ap"], values["f1"]] for label, values in metric["per_class"].items()])
            report.table(["Phát hiện", "Precision", "Recall", "Specificity", "Dương / Âm / Thiếu"],
                [[label, values["precision"], values["recall"], values["specificity"],
                  f"{values['positive']} / {values['negative']} / {values['missing']}"] for label, values in metric["per_class"].items()])
            for chart in ("auroc", "ap", "f1"):
                path = entry["root"] / "figures" / f"{chart}.png"
                if path.is_file():
                    report.image(path, f"{dataset}: {chart.upper()} theo lớp, chỉ hiển thị giá trị xác định")
            valid_index = next((i for i, (label, values) in enumerate(metric["per_class"].items()) if values["auroc"] is not None), None)
            if valid_index is not None:
                chart = entry["root"] / "figures" / f"class-{valid_index:03d}.png"
                if chart.is_file():
                    report.image(chart, f"{dataset}: ROC, precision recall và confusion matrix của lớp hợp lệ đầu tiên theo thứ tự lưu")
        elif number == 10:
            report.paragraph("Các số liệu công bố dưới đây được trích từ Supplementary Tables S9–S11 và S13. "
                "PadChest là zero shot trên 55 findings; Indiana trên 30 findings. CheXpert S13 là supervised linear probe "
                "trên 13 findings, nên chênh lệch với zero shot chỉ dùng làm bối cảnh. Macro tham chiếu được tính từ các AUROC "
                "đã làm tròn ba chữ số; không thay thế bằng giá trị headline xấp xỉ.")
            report.table(["Dữ liệu", "Thí nghiệm công bố", "Lớp", "Macro từ bảng làm tròn"],
                [[dataset, "Supervised linear probe" if dataset == "chexpert" else "Zero shot", len(references(dataset)),
                  float(references(dataset).auroc.mean())] for dataset in DATASETS])
            for dataset, entry in runs.items():
                if not entry["metrics"]:
                    continue
                comparison, summary = compare(entry["run"], entry["metrics"])
                report.heading(dataset, 3)
                report.paragraph("Khả năng so sánh: " + summary["reason"] + ". Chưa xác minh tái lập chính xác.")
                report.table(["Phát hiện", "Paper", "Thực nghiệm", "|Δ|", "KTC thực nghiệm"],
                    [[r["disease"], r["paper_auroc"], r["reproduced_auroc"], r["absolute_difference"],
                      ci_text({"low": r["ci_low"], "high": r["ci_high"]})] for r in comparison])
                report.paragraph("Support dương/âm, KTC công bố và trạng thái so sánh từng lớp nằm trong comparison.csv/json; "
                    "chỉ đối chiếu mô tả, không suy ra ý nghĩa thống kê của chênh lệch giữa hai nghiên cứu.")
        elif number == 11:
            report.paragraph("FP và FN được xác định theo nhãn tham chiếu và ngưỡng đã khai báo. Nhãn báo cáo có thể thiếu bệnh hoặc "
                "khác với hình ảnh cụ thể trong cùng nghiên cứu; do đó không coi các trường hợp này là lỗi lâm sàng đã xác minh. "
                "Cần bác sĩ đọc lại ảnh và báo cáo để phân tích nguyên nhân. Không suy diễn tổn thương từ tên tệp hay điểm mô hình.")
            for dataset, entry in runs.items():
                if not entry["metrics"]:
                    continue
                report.heading(dataset, 3)
                report.table(["Phát hiện", "TN", "FP", "FN", "TP"], [[label, m["tn"], m["fp"], m["fn"], m["tp"]]
                    for label, m in entry["metrics"]["per_class"].items()])
                path = entry["root"] / "errors.csv"
                if path.is_file():
                    examples = pd.read_csv(path).head(6)
                    if not examples.empty:
                        report.table(["Phát hiện", "Loại", "Sample ID", "Điểm"], examples[["disease", "error", "sample_id", "score"]].values.tolist())
                    else:
                        report.paragraph("Không có FP/FN trên các nhãn xác định tại ngưỡng đã dùng; kiểm tra support trước khi diễn giải.")
                report.paragraph("Danh sách tối đa 20 FP và 20 FN mỗi lớp được lưu tại errors.csv, kèm đường dẫn ảnh để kiểm tra trên máy chủ.")
        elif number == 12:
            report.paragraph("Các hạn chế gồm sự khác biệt cohort Kaggle với benchmark gốc; nhãn yếu ở cấp báo cáo; ảnh lateral/frontal; "
                "tương quan nhiều ảnh cùng bệnh nhân; thiếu bản đồ bệnh nhân Indiana; chính sách missing/uncertain; "
                "khác biệt chuyển ảnh 16 bit; recipe prompt SFR chưa được xác nhận đầy đủ; sai số số học theo thiết bị và chunk. "
                "Không có manifest công bố để tự động chứng minh toàn bộ cohort và row alignment giống paper; checksum chỉ xác nhận "
                "các tệp đã đăng ký không bị đổi. CheXpert train/valid có nguy cơ trùng nguồn CheXpert Plus trong pretraining. "
                "CheXpert zero shot không tái lập supervised linear probe. Không huấn luyện hay thay đổi pretrained weights.")
            report.paragraph("Không có mẫu báo cáo CXR LLaVA được cung cấp trong kho mã hiện tại. Báo cáo này theo 14 mục người dùng "
                "yêu cầu; có thể truyền template DOCX để tái sử dụng định dạng. Các kiểm tra synthetic không chứng minh độ đúng của "
                "nhãn thật hay khả năng chạy model trên phần cứng chưa được thử nghiệm.")
        elif number == 13:
            completed = sum(bool(entry["metrics"]) for entry in runs.values())
            report.paragraph(f"Có {completed}/3 bộ dữ liệu với kết quả hoàn tất được cung cấp cho báo cáo. "
                "Chưa thể kết luận tái lập chính xác chỉ từ tên mô hình hoặc một macro AUROC. "
                "Bước tiếp theo là xác minh cohort, nguồn nhãn và recipe SFR; mở rộng support; kiểm tra lỗi với bác sĩ; "
                "và đối chiếu trên cùng tập findings. Các thí nghiệm supervised tương lai nằm ngoài pipeline zero shot hiện tại.")
        else:
            report.paragraph("Tệp thực nghiệm gồm cohort.json, data_check.json, identity.json, run.json, prediction_parts, predictions.json/csv, "
                "metrics.json/csv, summary.csv, comparison.json/csv, errors.csv và figures. metrics.csv chứa KTC 95% cho AUROC, AP, "
                "precision, recall, specificity, F1 và các ô confusion matrix; không cần chạy lại mô hình để tính lại chỉ số.")
            for dataset, entry in runs.items():
                report.paragraph(f"{dataset}: run.json SHA256={file_hash(entry['root'] / 'run.json')}; đường dẫn={entry['root']}.")
            report.paragraph(f"[1] Han và cộng sự. CLEAR an auditable foundation model for radiology grounded in clinical concepts. "
                f"Nature Biomedical Engineering, 2026. {PAPER_URL}")
            report.paragraph(f"[2] Supplementary Information, Tables S1, S6, S9–S11 và S13. {SUPPLEMENT_URL}")
            report.paragraph("[3] Mã và assets chính thức: https://github.com/peterhan91/CLEAR ; https://huggingface.co/peterhan91/CLEAR")
            report.paragraph("[4] SFR Embedding Mistral: https://huggingface.co/Salesforce/SFR-Embedding-Mistral")
            report.paragraph("[5] Indiana: https://www.kaggle.com/datasets/raddar/chest-xrays-indiana-university ; "
                "PadChest sample: https://www.kaggle.com/datasets/raddar/padchest-chest-xrays-sample ; "
                "CheXpert: https://www.kaggle.com/datasets/ashery/chexpert")
    return report


def generate(paths, output_dir="reports", *, template=None, overwrite=False):
    root = Path(output_dir)
    root.mkdir(parents=True, exist_ok=True)
    markdown, word = root / "CLEAR_evaluation.md", root / "CLEAR_evaluation.docx"
    if not overwrite and (markdown.exists() or word.exists()):
        raise FileExistsError("Report already exists; choose another --output-dir or explicitly pass --overwrite")
    report = build_report(read_runs(paths))
    temporary_word = root / ".CLEAR_evaluation.tmp.docx"
    report.docx(temporary_word, template)
    temporary_markdown = root / ".CLEAR_evaluation.tmp.md"
    temporary_markdown.write_text(report.markdown(markdown), encoding="utf-8")
    temporary_word.replace(word)
    temporary_markdown.replace(markdown)
    return markdown, word
